"""Нарезка загружаемых файлов (PDF / PPTX), строки длинных пунктов, справочные статьи сценария."""

from pathlib import Path

import pytest

from app.kb.chunker import chunk_pages, load_all_docs, sub_lines
from app.kb.ingest import (
    DRAFT_PORTAL,
    SCENARIO_XLSX,
    _drop_questions,
    chunk_file,
    load_scenario_articles,
)

DOCS = Path(__file__).resolve().parents[2] / "docs"


def _pdf_with_text(path: Path, lines: list[str]) -> None:
    """Минимальный PDF с текстовым слоем (Helvetica, латиница) — без сторонних библиотек."""
    y, ops = 760, []
    for ln in lines:
        safe = ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"BT /F1 11 Tf 50 {y} Td ({safe}) Tj ET")
        y -= 16
    stream = "\n".join(ops).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def test_pdf_points_with_pages(tmp_path):
    p = tmp_path / "rules.pdf"
    _pdf_with_text(p, ["1. GENERAL RULES", "1.1. Work starts at 9:00.", "1.2. Lunch break is 48 minutes.",
                       "2. VACATION", "2.1. Vacation is 28 days."])
    chunks = chunk_file(p, "Test rules")
    assert [c.point_no for c in chunks] == ["1.1", "1.2", "2.1"]
    assert chunks[1].text == "1.2. Lunch break is 48 minutes."
    assert chunks[1].source_label == "Test rules, п. 1.2, стр. 1"
    assert chunks[2].path.startswith("Раздел 2")


def test_pdf_without_text_layer_is_rejected(tmp_path):
    from PIL import Image

    p = tmp_path / "scan.pdf"
    Image.new("RGB", (100, 100), "white").save(p, "PDF")
    with pytest.raises(ValueError, match="текстового слоя"):
        chunk_file(p, "Скан")


def test_pdf_pages_without_points_are_paragraphs():
    pages = [(1, "Памятка новичку.\n\nПропуск выдают на проходной."), (2, "Обед в столовой на 2 этаже.")]
    chunks = chunk_pages(pages, "doc", "Памятка")
    assert [c.source_label for c in chunks] == ["Памятка, стр. 1", "Памятка, стр. 2"]
    assert "Пропуск выдают на проходной." in chunks[0].text


def test_pptx_slide_is_chunk(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Деловые подарки"
    s1.placeholders[1].text = "Подарок дороже 3000 рублей принимать нельзя."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Горячая линия"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box.text_frame.text = "Телефон доверия: 8 800 000-00-00 (синтетика)"
    p = tmp_path / "policy.pptx"
    prs.save(p)

    chunks = chunk_file(p, "Политика (тест)")
    assert [c.source_label for c in chunks] == ["Политика (тест), слайд 1", "Политика (тест), слайд 2"]
    assert chunks[0].text.startswith("Деловые подарки")
    assert "3000 рублей" in chunks[0].text
    assert "Телефон доверия" in chunks[1].text


def test_unsupported_format(tmp_path):
    p = tmp_path / "x.docx"
    p.write_bytes(b"x")
    with pytest.raises(ValueError, match="не поддерживается"):
        chunk_file(p, "x")


def test_sub_lines_only_for_long_points():
    chunks = {c.key: c for c in load_all_docs(DOCS)}
    lines = sub_lines(chunks["pvtr:6.4.2"])
    assert any("курение" in ln for _, ln in lines)
    assert all(ln in chunks["pvtr:6.4.2"].text for _, ln in lines)  # дословно
    assert sub_lines(chunks["pvtr:6.6.1"]) == []


def test_drop_questions_keeps_text_verbatim():
    t = "Чтобы дорога до офиса была комфортнее, у нас есть корпоративный транспорт! Рассказать о маршрутах?"
    assert _drop_questions(t) == "Чтобы дорога до офиса была комфортнее, у нас есть корпоративный транспорт!"


def test_draft_article_is_marked_synthetic():
    assert "синтетика" in DRAFT_PORTAL.source


@pytest.mark.skipif(not SCENARIO_XLSX.exists(), reason="нет xlsx сценария (data/raw не в git)")
def test_scenario_articles():
    arts = {(a.slug, tuple(a.site_codes)): a for a in load_scenario_articles()}
    docs_gtp = arts[("documents", ("GTP",))]
    assert "СНИЛС" in docs_gtp.text and "Фотографии 3*4" in docs_gtp.text
    assert "Фотографии" not in arts[("documents", ("GrPP", "Luzhniki"))].text
    assert not arts[("transport", ("GrPP",))].text.endswith("?")
    assert "address-book" in arts[("directory", ("GrPP", "Luzhniki"))].text
    assert "/Food" in arts[("food", ())].text
    assert ("gym", ("GTP",)) in arts
    keys = [a.key for a in arts.values()]
    assert len(keys) == len(set(keys))
    assert all(len(k) <= 64 for k in keys)
