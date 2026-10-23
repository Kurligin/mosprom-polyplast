"""Наполнение базы знаний: документы, справочные статьи сценария, ответы HR → kb_chunks с эмбеддингами.

CLI (в Docker): python -m app.kb.ingest --all            (make kb)
                python -m app.kb.ingest --file x.pdf --title "…" [--kt] [--sites GTP,GrPP]

Функции не коммитят: коммит делает вызывающий код (CLI, API).
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import HrAnswer, KbChunk, KbDocument
from app.kb import embed as E
from app.kb.chunker import (
    DOCS,
    Chunk,
    chunk_document,
    chunk_markdown,
    chunk_pages,
    chunk_slides,
    slides_to_md,
    sub_lines,
)

log = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parents[2]
DOCS_DIR = ROOT / "docs"
SCENARIO_XLSX = ROOT / "data" / "raw" / "korsik_scenario.xlsx"

# Человеческие названия документов для панели HR (H1 в md есть не у всех)
DOC_TITLES = {
    "pvtr": "Правила внутреннего трудового распорядка",
    "anticorruption-policy": "Антикоррупционная политика",
    "trade-secret": "Положение о коммерческой тайне и конфиденциальной информации",
    "quality-ecology-policy": "Политика в области качества и экологии",
}

SCENARIO_TYPE = "scenario"
DRAFT_LABEL = "Черновик HR, ждёт подтверждения"


@dataclass
class Article:
    """Справочная статья: дословные фразы из шагов сценария (или черновик HR)."""

    slug: str
    title: str
    text: str
    site_codes: list[str] = field(default_factory=list)
    source: str = ""
    kind: str = "Памятка Корсика"  # подпись источника: «Памятка Корсика «…»» или «Статья HR «…»»

    @property
    def key(self) -> str:
        return f"scenario:{self.slug}" + (f"@{'+'.join(self.site_codes)}" if self.site_codes else "")

    @property
    def label(self) -> str:
        return f"{self.kind} «{self.title}»"


# Синтетический черновик для демо: статья HR, которую ещё не подтвердили
DRAFT_PORTAL = Article(
    slug="portal-access",
    title="Доступ к корпоративному порталу",
    text=(
        "Как получить доступ к корпоративному порталу my.polyplastic.ru:\n"
        "1. В первый рабочий день IT заводит учётную запись. Логин и временный пароль выдают вместе с техникой.\n"
        "2. Портал открывается с рабочего компьютера в корпоративной сети по адресу https://my.polyplastic.ru.\n"
        "3. При первом входе смените временный пароль. Пароль никому не сообщайте.\n"
        "4. С личного устройства портал доступен только через корпоративную учётную запись.\n"
        "5. Если войти не получается, напишите руководителю или в IT-поддержку: логин и текст ошибки."
    ),
    source="синтетика: демо-черновик статьи HR",
    kind="Статья HR",
)


# --- справочные статьи из сценария -------------------------------------------

_SENT_RE = re.compile(r"(?<=[.!?])\s+")


def _sentences(text: str) -> list[str]:
    return [s for s in _SENT_RE.split(" ".join(text.split())) if s]


def _drop_questions(text: str) -> str:
    """Убрать вопросы в конце («Рассказать подробнее?»): в справке они не нужны. Остальное дословно."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    out = []
    for p in paras:
        sents = _sentences(p)
        while sents and sents[-1].endswith("?"):
            sents.pop()
        if sents:
            out.append(" ".join(sents))
    return "\n\n".join(out)


def _note(comment: str | None, start: str) -> str:
    """Список из «Примечания» шага (там он дословно, в сообщении — картинкой)."""
    if not comment:
        return ""
    i = comment.find(start)
    if i < 0:
        return ""
    return comment[i:].split(" | ")[0].strip()


def _reply(step, word: str) -> str:
    for b in step.buttons or []:
        if b.reply and b.label.lower().startswith(word):
            return b.reply.strip()
    return ""


def _sentence_with(text: str, needle: str) -> str:
    for s in _sentences(text):
        if needle in s:
            return s
    return ""


def _cap(s: str) -> str:
    return s[:1].upper() + s[1:] if s else s


def scenario_articles(steps) -> list[Article]:
    """Справочные статьи из шагов сценария: документы, транспорт, доступы, обед, справочник и т. п.

    Текст — дословные фразы шага; вопросы-приглашения в конце («Интересно?») убраны.
    Шаги-приветствия, опросы и поздравления в базу не идут.
    """
    arts: list[Article] = []

    def add(slug, title, text, step):
        text = (text or "").strip()
        if text and not any(a.slug == slug and a.site_codes == list(step.site_tags) for a in arts):
            arts.append(Article(slug, title, text, list(step.site_tags), f"сценарий, шаг {step.number}"))

    for s in steps:
        topic = (s.topic or "").lower()
        text = s.text or ""
        low = text.lower()
        if s.code == "1.5" and "документ" in topic:
            note = _note(s.comment, "Необходимые документы")
            add("documents", "Документы для оформления", f"{text}\n\n{note}" if note else text, s)
        elif "корпоративном транспорте" in topic:
            title = "Как добраться до офиса" if "добраться" in low else "Корпоративный транспорт"
            add("transport", title, _drop_questions(text), s)
        elif s.code == "3.1":
            paras = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
            body = _drop_questions("\n\n".join(paras[1:]) if len(paras) > 1 else text)
            no = _reply(s, "нет")
            add("access", "Доступы и техника в первый день", "\n\n".join(x for x in (body, no) if x), s)
        elif s.code == "3.2":
            food = _sentence_with(_reply(s, "нет"), "/Food")
            # «Если ты передумаешь, ты всегда можешь…» → «Ты всегда можешь…»
            food = _cap(re.sub(r"^Если ты передумаешь,\s*", "", food))
            add("food", "Где пообедать", food, s)
        elif s.code == "3.3" and "справочник" in low:
            add("directory", "Справочник контактов сотрудников", text, s)
        elif s.code == "3.4":
            add("mission", "Миссия компании", text, s)
        elif "спортивном зале" in topic:
            add("gym", "Спортивный комплекс", _drop_questions(text), s)
        elif any("адаптационный план" in d.lower() for d in s.directives or []):
            add("adaptation-plan", "Адаптационный план", text, s)
        elif low.startswith("как мы проводим корпоративные мероприятия"):
            add("events", "Корпоративные мероприятия", text, s)
        elif "фактами о своей организации" in low:
            add("facts", "Интересные факты о компании", text, s)
        elif "/HR" in text:
            add("contact-hr", "Как связаться с HR", _sentence_with(text, "/HR"), s)
    return arts


def load_scenario_articles(path: Path = SCENARIO_XLSX) -> list[Article]:
    """Статьи из xlsx сценария (только чтение). Нет файла — пустой список."""
    if not path.exists():
        log.warning("Сценарий %s не найден: справочные статьи пропущены", path)
        return []
    from app.scenario.importer import parse_xlsx

    return scenario_articles(parse_xlsx(path))


# --- запись в БД -------------------------------------------------------------


def _embed_input(title: str, path: str, text: str) -> str:
    return f"{title}. {path}\n{text}"


@dataclass
class _Spec:
    """Фрагмент к записи: ключ, поля строки kb_chunks и текст для эмбеддинга."""

    key: str
    path: str
    point_no: str | None
    source_label: str
    page: int | None
    text: str
    is_kt: bool
    embed_input: str


async def _open_document(session: AsyncSession, *, title: str, doc_type: str, source: str | None, match_by: str,
                         is_kt: bool, site_codes: list[str], status: str,
                         uploaded_by: int | None) -> tuple[KbDocument, bool]:
    """Найти документ (по названию или источнику) или создать. → (документ, можно ли брать старые векторы)."""
    col = KbDocument.title if match_by == "title" else KbDocument.source
    val = title if match_by == "title" else source
    doc = (await session.execute(select(KbDocument).where(col == val, KbDocument.doc_type != "hr_answer")
                                 .order_by(KbDocument.id))).scalars().first()
    if doc is None:
        doc = KbDocument(title=title, doc_type=doc_type, source=source, site_codes=site_codes, is_kt=is_kt,
                         status=status, version=1, uploaded_by=uploaded_by)
        session.add(doc)
        await session.flush()
        return doc, False
    same_title = doc.title == title
    meta_changed = (doc.title, doc.doc_type, doc.source, list(doc.site_codes or []), doc.is_kt, doc.status) != (
        title, doc_type, source, list(site_codes), is_kt, status)
    doc.title, doc.doc_type, doc.source = title, doc_type, source
    doc.site_codes, doc.is_kt, doc.status = list(site_codes), is_kt, status
    if uploaded_by is not None:
        doc.uploaded_by = uploaded_by
    doc._kb_meta_changed = meta_changed  # type: ignore[attr-defined] — для счётчика версий
    return doc, same_title


async def _sync_chunks(session: AsyncSession, doc: KbDocument, specs: list[_Spec], reuse_vectors: bool) -> None:
    """Привести фрагменты документа к specs: строки с тем же ключом обновляются на месте (id сохраняется,
    ссылки из questions_log не теряются), новые добавляются, лишние удаляются. Векторы неизменённых
    фрагментов не пересчитываем. Документ, у которого что-то поменялось, получает версию +1."""
    old: dict[str | None, KbChunk] = {}
    for r in (await session.execute(select(KbChunk).where(KbChunk.document_id == doc.id)
                                    .order_by(KbChunk.id))).scalars().all():
        if r.chunk_key in old:
            await session.delete(r)  # повтор ключа (старые версии нарезки)
        else:
            old[r.chunk_key] = r
    seen: dict[str, int] = {}
    for sp in specs:  # ключ должен быть уникален внутри документа
        seen[sp.key] = seen.get(sp.key, 0) + 1
        if seen[sp.key] > 1:
            sp.key = f"{sp.key}#{seen[sp.key]}"[:64]
    rows: list[KbChunk] = []
    todo: list[int] = []
    changed = bool(getattr(doc, "_kb_meta_changed", False)) or set(old) != {sp.key for sp in specs}
    for i, sp in enumerate(specs):
        row = old.get(sp.key)
        same_text = row is not None and row.text == sp.text and row.path == sp.path
        if row is None:
            row = KbChunk(document_id=doc.id, chunk_key=sp.key)
            session.add(row)
        elif not same_text:
            row.tsv = None  # триггер пересчитает tsvector по новому тексту
        if not same_text or row.embedding is None or not reuse_vectors:
            todo.append(i)
            changed = True
        row.path, row.point_no, row.source_label = sp.path, sp.point_no, sp.source_label
        row.page, row.text, row.is_kt = sp.page, sp.text, sp.is_kt
        rows.append(row)
    keep = {sp.key for sp in specs}
    for key, row in old.items():
        if key not in keep:
            await session.delete(row)
    vecs = await E.embed([specs[i].embed_input for i in todo])
    for i, v in zip(todo, vecs, strict=True):
        rows[i].embedding = v
    if changed and old:
        doc.version = (doc.version or 1) + 1
    await session.flush()


def _specs(doc: KbDocument, chunks: list[Chunk], key_prefix: str | None = None, split_long: bool = True) -> list[_Spec]:
    """Фрагменты документа + строки длинных пунктов (отдельно ищутся, в ответе — пункт с выдержкой)."""
    out: list[_Spec] = []
    for c in chunks:
        key = (c.key if key_prefix is None else f"{key_prefix}:{c.point_no or c.page}")[:64]
        out.append(_Spec(key, c.path, c.point_no, c.source_label, c.page, c.text, c.is_kt,
                         _embed_input(doc.title, c.path, c.text)))
        if not split_long:
            continue
        intro = c.text.splitlines()[0].strip() if c.text else ""
        for i, line in sub_lines(c):
            out.append(_Spec(f"{key}/{i}"[:64], c.path, c.point_no, c.source_label, c.page, line, c.is_kt,
                             _embed_input(doc.title, c.path, f"{intro}\n{line}")))
    return out


async def ingest_markdown_doc(session: AsyncSession, path: Path) -> tuple[KbDocument, int]:
    """Документ из docs/ (ПВТР, презентации) — известная нарезка chunker."""
    slug = path.stem
    _short, _mode, is_kt = DOCS[slug]
    chunks = chunk_document(path)
    doc, reuse = await _open_document(session, title=DOC_TITLES.get(slug, chunks[0].doc_title if chunks else slug),
                                      doc_type="md", source=f"docs/{path.name}", match_by="source", is_kt=is_kt,
                                      site_codes=[], status="confirmed", uploaded_by=None)
    await _sync_chunks(session, doc, _specs(doc, chunks), reuse)
    return doc, len(chunks)


async def ingest_article(session: AsyncSession, art: Article, *, status: str | None = "confirmed",
                         doc_type: str = SCENARIO_TYPE) -> tuple[KbDocument, int]:
    """Статья = документ из одного фрагмента. status=None — оставить статус как есть (новая — черновик)."""
    source = f"{art.key} | {art.source}"  # ключ статьи в source: повторный импорт находит ту же запись
    if status is None:
        old = (await session.execute(select(KbDocument.status).where(KbDocument.source == source))).scalars().first()
        status = old or "draft"
    title = art.title + (f" ({', '.join(art.site_codes)})" if art.site_codes else "")
    doc, reuse = await _open_document(session, title=title, doc_type=doc_type, source=source, match_by="source",
                                      is_kt=False, site_codes=art.site_codes, status=status, uploaded_by=None)
    spec = _Spec(art.key[:64], art.title, None, art.label, None, art.text, False,
                 _embed_input(art.title, "", art.text))
    await _sync_chunks(session, doc, [spec], reuse)
    return doc, 1


async def ingest_all(session: AsyncSession, docs_dir: Path = DOCS_DIR, scenario_xlsx: Path = SCENARIO_XLSX) -> dict:
    """docs/*.md + справочные статьи сценария + черновик «Доступ к корпоративному порталу»."""
    report: dict = {"documents": {}, "articles": [], "draft": None}
    for slug in DOCS:
        p = docs_dir / f"{slug}.md"
        if p.exists():
            doc, n = await ingest_markdown_doc(session, p)
            report["documents"][doc.title] = n

    arts = load_scenario_articles(scenario_xlsx)
    keep: set[str] = set()
    for art in arts:
        doc, _ = await ingest_article(session, art)
        keep.add(doc.source or "")
        report["articles"].append(doc.title)
    # статьи, которых больше нет в сценарии, убираем
    old = (await session.execute(select(KbDocument).where(KbDocument.doc_type == SCENARIO_TYPE))).scalars().all()
    for d in old:
        if d.source not in keep:
            await session.delete(d)

    # черновик HR: если HR уже подтвердил статью в панели, повторный импорт статус не сбрасывает
    doc, _ = await ingest_article(session, DRAFT_PORTAL, status=None, doc_type="article")
    report["draft"] = doc.title
    await session.flush()
    return report


# --- загрузка файлов HR -------------------------------------------------------


def _pdf_pages(path: Path) -> list[tuple[int, str]]:
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        return [(i, page.extract_text() or "") for i, page in enumerate(pdf.pages, start=1)]


def _shape_texts(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # группа
        for s in shape.shapes:
            out.extend(_shape_texts(s))
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = " ".join(p.text.split())
            if t:
                out.append(t)
    return out


def _pptx_slides(path: Path) -> list[tuple[int, str, str]]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = []
    for no, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        body: list[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            body.extend(_shape_texts(shape))
        slides.append((no, " ".join(title.split()), "\n\n".join(body)))
    return slides


def chunk_file(path: Path, title: str, is_kt: bool = False) -> list[Chunk]:
    """PDF (страницы с номерами), PPTX (слайд = фрагмент) или MD → фрагменты с подписями."""
    suffix = path.suffix.lower()
    short = title
    if suffix == ".pdf":
        pages = _pdf_pages(path)
        if not any(t.strip() for _, t in pages):
            raise ValueError("В PDF нет текстового слоя (скан). Загрузите PDF с текстом.")
        return chunk_pages(pages, "doc", short, is_kt)
    if suffix == ".pptx":
        return chunk_slides(slides_to_md(_pptx_slides(path)), "doc", short, is_kt)
    if suffix in (".md", ".txt"):
        return chunk_markdown(path.read_text(encoding="utf-8"), "doc", short, is_kt)
    raise ValueError(f"Формат {suffix or 'без расширения'} не поддерживается: нужен PDF или PPTX.")


async def ingest_file(session: AsyncSession, path: str | Path, *, title: str, is_kt: bool = False,
                      site_codes: list[str] | None = None, uploaded_by: int | None = None,
                      status: str = "confirmed") -> tuple[KbDocument, int]:
    """Документ HR (PDF / PPTX / MD) → фрагменты. Документ с тем же названием заменяется (версия +1).

    ValueError — формат не поддерживается, в PDF нет текста или фрагментов не получилось.
    """
    path = Path(path)
    title = " ".join((title or path.stem).split())
    chunks = chunk_file(path, title, is_kt)
    if not chunks:
        raise ValueError("Не удалось выделить текст из файла.")
    doc_type = {".pdf": "pdf", ".pptx": "pptx"}.get(path.suffix.lower(), "md")
    doc, reuse = await _open_document(session, title=title, doc_type=doc_type, source=path.name, match_by="title",
                                      is_kt=is_kt, site_codes=list(site_codes or []), status=status,
                                      uploaded_by=uploaded_by)
    await _sync_chunks(session, doc, _specs(doc, chunks, key_prefix=f"doc{doc.id}"), reuse)
    return doc, len(chunks)


# --- ответы HR ------------------------------------------------------------------


def hr_answer_label(hr_answer: HrAnswer) -> str:
    when = hr_answer.created_at
    return f"Ответ HR от {when:%d.%m.%Y}" if when else "Ответ HR"


async def index_hr_answer(session: AsyncSession, hr_answer: HrAnswer) -> KbChunk:
    """Ответ HR → фрагмент базы (создать или обновить). В поиск попадает только при status == "confirmed"."""
    if hr_answer.id is None:
        session.add(hr_answer)
        await session.flush()
    chunk = (await session.execute(select(KbChunk).where(KbChunk.hr_answer_id == hr_answer.id))).scalars().first()
    question = " ".join((hr_answer.question or "").split())
    if chunk is None:
        chunk = KbChunk(hr_answer_id=hr_answer.id, is_kt=False)
        session.add(chunk)
    chunk.path = question
    chunk.text = hr_answer.answer
    chunk.chunk_key = f"hr:{hr_answer.id}"
    chunk.source_label = hr_answer_label(hr_answer)
    chunk.tsv = None  # триггер пересчитает по новому вопросу и ответу
    chunk.embedding = (await E.embed([f"{question}\n{hr_answer.answer}"]))[0]
    await session.flush()
    return chunk


# --- CLI ----------------------------------------------------------------------------


async def _main(args) -> None:
    from app.db.session import async_session

    async with async_session() as session:
        if args.all:
            rep = await ingest_all(session)
            await session.commit()
            print("Документы:")
            for t, n in rep["documents"].items():
                print(f"  {t}: {n} пунктов")
            print(f"Справочные статьи сценария: {len(rep['articles'])}")
            for t in rep["articles"]:
                print(f"  {t}")
            print(f"Черновик HR: {rep['draft']}")
        if args.file:
            doc, n = await ingest_file(session, args.file, title=args.title or Path(args.file).stem, is_kt=args.kt,
                                       site_codes=[s for s in (args.sites or "").split(",") if s])
            await session.commit()
            print(f"{doc.title}: {n} фрагментов, версия {doc.version}")
        total = (await session.execute(select(KbChunk.id))).scalars().all()
        print(f"Всего строк в kb_chunks: {len(total)}")


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(name)s %(message)s")
    ap = argparse.ArgumentParser(description="Наполнение базы знаний Корсика")
    ap.add_argument("--all", action="store_true", help="docs/*.md, статьи сценария, черновик HR")
    ap.add_argument("--file", help="PDF / PPTX / MD документа HR")
    ap.add_argument("--title")
    ap.add_argument("--kt", action="store_true", help="коммерческая тайна")
    ap.add_argument("--sites", help="коды площадок через запятую; пусто — все")
    args = ap.parse_args()
    if not args.all and not args.file:
        ap.error("укажите --all или --file")
    asyncio.run(_main(args))


if __name__ == "__main__":
    main()
